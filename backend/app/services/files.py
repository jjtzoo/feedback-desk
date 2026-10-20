"""Attachments (brief §4.1, §8; D10).

PDF, Word (.docx), Excel (.xlsx), PNG, JPG and WEBP; max 10 MB each, 5 per item. The type is
checked by content, not only by extension, and macro-enabled Office files are rejected.
Word and Excel are turned into text on upload. Files live in a private bucket (or a local
folder in development) and are served through 15-minute signed links.
"""

import asyncio
import hashlib
import hmac
import io
import time
import uuid
import zipfile
from dataclasses import dataclass
from datetime import timedelta
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.errors import NotFound, ServiceError
from app.models import Attachment

ALLOWED = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
MAX_EXTRACT_CHARS = 20_000


class FileRejected(ServiceError):
    status_code = 415

    def __init__(self, message: str) -> None:
        super().__init__("file_rejected", message)


def sniff(content: bytes, file_name: str) -> str:
    """The MIME type from the bytes themselves; it must agree with the extension."""
    ext = Path(file_name).suffix.lower()
    if ext in (".docm", ".xlsm", ".dotm", ".xltm"):
        raise FileRejected("Macro-enabled Office files can't be uploaded.")
    if ext not in ALLOWED:
        raise FileRejected("Use PDF, Word (.docx), Excel (.xlsx), PNG, JPG or WEBP.")
    if content.startswith(b"%PDF-"):
        found = "application/pdf"
    elif content.startswith(b"\x89PNG\r\n\x1a\n"):
        found = "image/png"
    elif content.startswith(b"\xff\xd8\xff"):
        found = "image/jpeg"
    elif content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        found = "image/webp"
    elif content.startswith(b"PK\x03\x04"):
        found = _office_type(content)
    else:
        raise FileRejected("That file's contents don't match a supported type.")
    if found != ALLOWED[ext]:
        raise FileRejected("That file's contents don't match its extension.")
    return found


def _office_type(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            names = set(z.namelist())
            types = z.read("[Content_Types].xml").decode("utf-8", "replace") if "[Content_Types].xml" in names else ""
    except zipfile.BadZipFile as exc:
        raise FileRejected("That Office file is damaged.") from exc
    if "word/vbaProject.bin" in names or "xl/vbaProject.bin" in names or "macroEnabled" in types:
        raise FileRejected("Macro-enabled Office files can't be uploaded.")
    if "word/document.xml" in names:
        return ALLOWED[".docx"]
    if "xl/workbook.xml" in names:
        return ALLOWED[".xlsx"]
    raise FileRejected("That file's contents don't match a supported type.")


def extract_text(content: bytes, mime: str) -> str | None:
    """Text from Word and Excel files, for the agent and the developer brief."""
    if mime == ALLOWED[".docx"]:
        import docx

        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)[:MAX_EXTRACT_CHARS]
    if mime == ALLOWED[".xlsx"]:
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in book.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    lines.append(" | ".join(cells))
                if sum(len(line) for line in lines) > MAX_EXTRACT_CHARS:
                    break
        book.close()
        return "\n".join(lines)[:MAX_EXTRACT_CHARS]
    return None


# --- storage ------------------------------------------------------------------------


def _local_path(key: str) -> Path:
    return Path(get_settings().storage_dir) / key


async def _put(key: str, content: bytes, mime: str) -> None:
    s = get_settings()
    if s.storage_backend == "gcs":

        def upload() -> None:
            from google.cloud import storage

            storage.Client().bucket(s.gcs_bucket).blob(key).upload_from_string(content, content_type=mime)

        await asyncio.to_thread(upload)
        return
    path = _local_path(key)
    path.parent.mkdir(parents=True, exist_ok=True)
    await asyncio.to_thread(path.write_bytes, content)


async def read(att: Attachment) -> bytes:
    s = get_settings()
    if s.storage_backend == "gcs":

        def download() -> bytes:
            from google.cloud import storage

            data: bytes = storage.Client().bucket(s.gcs_bucket).blob(att.storage_path).download_as_bytes()
            return data

        return await asyncio.to_thread(download)
    return await asyncio.to_thread(_local_path(att.storage_path).read_bytes)


@dataclass
class Upload:
    file_name: str
    content: bytes


async def save(
    db: AsyncSession,
    upload: Upload,
    uploaded_by: str,
    *,
    session_id: str | None = None,
    feedback_id: uuid.UUID | None = None,
) -> Attachment:
    s = get_settings()
    if len(upload.content) > s.max_file_mb * 1024 * 1024:
        raise FileRejected(f"Files can be up to {s.max_file_mb} MB.")
    if not upload.content:
        raise FileRejected("That file is empty.")
    owner = Attachment.session_id == session_id if session_id else Attachment.feedback_id == feedback_id
    count = int(await db.scalar(select(func.count()).select_from(Attachment).where(owner)) or 0)
    if count >= s.max_files_per_item:
        raise FileRejected(f"Up to {s.max_files_per_item} files per feedback.")
    mime = sniff(upload.content, upload.file_name)
    ext = Path(upload.file_name).suffix.lower()
    key = f"{time.strftime('%Y/%m')}/{uuid.uuid4().hex}{ext}"
    await _put(key, upload.content, mime)
    att = Attachment(
        feedback_id=feedback_id,
        session_id=session_id,
        uploaded_by=uploaded_by.lower(),
        file_name=Path(upload.file_name).name[:255],
        storage_path=key,
        mime_type=mime,
        size_bytes=len(upload.content),
        extracted_text=await asyncio.to_thread(extract_text, upload.content, mime),
    )
    db.add(att)
    await db.flush()
    return att


# --- signed links -------------------------------------------------------------------


def _sig(att_id: uuid.UUID, expires: int) -> str:
    key = get_settings().secret_key.encode()
    return hmac.new(key, f"{att_id}:{expires}".encode(), hashlib.sha256).hexdigest()


async def signed_url(att: Attachment) -> str:
    """A link valid for 15 minutes. Callers check the user may see the item first."""
    s = get_settings()
    if s.storage_backend == "gcs":

        def sign() -> str:
            from google.cloud import storage

            blob = storage.Client().bucket(s.gcs_bucket).blob(att.storage_path)
            url: str = blob.generate_signed_url(
                version="v4", expiration=timedelta(minutes=s.signed_url_minutes), method="GET"
            )
            return url

        return await asyncio.to_thread(sign)
    expires = int(time.time()) + s.signed_url_minutes * 60
    return f"/api/files/{att.id}/content?e={expires}&s={_sig(att.id, expires)}"


def check_local_signature(att_id: uuid.UUID, expires: int, signature: str) -> None:
    if expires < time.time() or not hmac.compare_digest(_sig(att_id, expires), signature):
        raise NotFound("That link has expired.")
