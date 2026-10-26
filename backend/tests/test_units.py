"""Unit tests: privacy masking, file checks, routing rules, draft guards, tool guards and model credentials."""

import io
import os
import zipfile
from dataclasses import dataclass, field
from typing import Any

import pytest
from sqlalchemy import select, update

from agent import callbacks, model, tools
from app import agent_runtime
from app.config import get_settings
from app.db import session_scope
from app.models import DigitalStaff, Feedback, Project, RoutingRule
from app.services import files, privacy
from app.services.drafts import apply_update, card, draft_hash, impact_chips, missing, quick_replies
from app.services.router import route

# --- privacy -------------------------------------------------------------------------------


def test_masks_passwords_cards_and_tfns_but_keeps_ordinary_numbers() -> None:
    text = "My password is Hunter2! card 4111 1111 1111 1111, TFN 123 456 782, invoice 1234567, call 0412345678"
    out = privacy.mask(text)
    assert "Hunter2" not in out and "[password removed]" in out
    assert "4111" not in out and "[card number removed]" in out
    assert "123 456 782" not in out and "[TFN removed]" in out
    assert "1234567" in out and "0412345678" in out


def test_masking_twice_changes_nothing() -> None:
    once = privacy.mask("My password is Hunter2, pin 4321, card 4111 1111 1111 1111, TFN 123 456 782")
    assert privacy.mask(once) == once and once.count("[password removed]") == 2


def test_abuse_about_a_named_colleague_is_spotted() -> None:
    assert privacy.is_abusive_about("Darlene is useless and lazy", ["Darlene", "Mario"])
    assert not privacy.is_abusive_about("This report is useless", ["Darlene"])
    assert not privacy.is_abusive_about("Darlene fixed it, great", ["Darlene"])


# --- files ---------------------------------------------------------------------------------

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def _docx(macro: bool = False) -> bytes:
    import docx

    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("Error 500 when saving the leave report")
    d.save(buf)
    if not macro:
        return buf.getvalue()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("word/vbaProject.bin", b"macro")
    return out.getvalue()


def _xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Staff", "CPD points"])  # type: ignore[union-attr]
    ws.append(["Sam", 12])  # type: ignore[union-attr]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_file_type_is_checked_by_content() -> None:
    assert files.sniff(PNG, "screen.png") == "image/png"
    assert files.sniff(b"%PDF-1.7 ...", "report.pdf") == "application/pdf"
    with pytest.raises(files.FileRejected):
        files.sniff(PNG, "report.pdf")  # extension says PDF, content says PNG
    with pytest.raises(files.FileRejected):
        files.sniff(b"MZ\x90\x00", "tool.exe")
    with pytest.raises(files.FileRejected):
        files.sniff(_docx(), "notes.docm")


def test_macro_enabled_office_files_are_rejected() -> None:
    with pytest.raises(files.FileRejected, match="Macro"):
        files.sniff(_docx(macro=True), "notes.docx")


def test_word_and_excel_become_text() -> None:
    assert "Error 500" in (files.extract_text(_docx(), files.ALLOWED[".docx"]) or "")
    text = files.extract_text(_xlsx(), files.ALLOWED[".xlsx"]) or ""
    assert "CPD points" in text and "Sam | 12" in text


# --- routing (brief §5.3) ------------------------------------------------------------------


async def _route(
    category: str, slug: str | None, confidence: float = 0.9, flagged: bool = False
) -> tuple[str | None, str]:
    async with session_scope() as db:
        project = await db.scalar(select(Project).where(Project.slug == slug)) if slug else None
        fb = Feedback(category=category, flagged=flagged)
        r = await route(db, fb, project, confidence)
        return (r.assignee.name if r.assignee else None), r.status


async def test_routes_to_the_project_owner_then_waits_for_approval() -> None:
    assert await _route("bug", "rake-agent") == ("Darlene", "awaiting_approval")
    assert await _route("feature_request", "cpd-tracker") == ("Mario", "awaiting_approval")


async def test_rule_overrides_come_first() -> None:
    assert await _route("access_request", "rake-agent") == ("Chris", "awaiting_approval")


async def test_low_confidence_or_general_goes_to_needs_review() -> None:
    assert await _route("bug", "rake-agent", confidence=0.5) == (None, "needs_review")
    assert await _route("bug", None) == (None, "needs_review")


async def test_away_owner_goes_to_needs_review() -> None:
    async with session_scope() as db:
        await db.execute(update(DigitalStaff).where(DigitalStaff.name == "Darlene").values(is_away=True))
    assert await _route("bug", "rake-agent") == (None, "needs_review")


async def test_praise_is_acknowledged_and_flagged_goes_to_admins() -> None:
    assert await _route("praise", "rake-agent") == ("Darlene", "acknowledged")
    assert await _route("bug", "rake-agent", flagged=True) == (None, "needs_review")


async def test_group_rule_picks_the_member_with_fewest_open_items() -> None:
    async with session_scope() as db:
        mario = await db.scalar(select(DigitalStaff).where(DigitalStaff.name == "Mario"))
        darlene = await db.scalar(select(DigitalStaff).where(DigitalStaff.name == "Darlene"))
        assert mario and darlene
        for s in (mario, darlene):
            db.add(RoutingRule(category="data_issue", assignee_id=s.id, priority_order=5))
        db.add(
            Feedback(
                submitter_email="x@redadair.example",
                submitter_name="X",
                title="t",
                summary="s",
                status="in_progress",
                assignee_id=darlene.id,
            )
        )
    assert (await _route("data_issue", "rake-agent"))[0] == "Mario"


# --- drafts and tool guards -------------------------------------------------------------------


COMPLETE: dict[str, Any] = {
    "project_slug": "rake-agent",
    "work_type": "fix",
    "category": "bug",
    "priority": "P2",
    "summary": "RAKE gives a blank answer",
    "title": "Blank answers",
}


def test_draft_validation_rejects_unknown_values() -> None:
    _, errors = apply_update(
        {},
        {"project_slug": "nope", "category": "pricing", "priority": "urgent", "project_confidence": "x"},
        {"rake-agent"},
    )
    assert len(errors) == 4
    draft, errors = apply_update(
        {}, {**COMPLETE, "steps": ["Open RAKE", "Ask anything"], "project_confidence": 0.9}, {"rake-agent"}
    )
    assert not errors and not missing(draft)
    assert draft["brief"]["steps"] == ["Open RAKE", "Ask anything"] and draft["confidence"]["project"] == 0.9


def test_a_tool_only_counts_as_named_when_the_person_said_it() -> None:
    from types import SimpleNamespace

    from app.services.catalog import mentions

    intranet = SimpleNamespace(
        name="Intranet site itself", slug="intranet", aliases=["intranet"], keywords=["home page", "tile", "menu"]
    )
    assert not mentions("page is slow to load", intranet)  # "page" alone names nothing
    assert mentions("The intranet home page is slow", intranet)
    assert mentions("the Flipbook tile on the home page goes to a 404", intranet)
    rake = SimpleNamespace(name="RAKE Agent", slug="rake-agent", aliases=["rake"], keywords=["policy answer"])
    assert mentions("RAKE gives a blank answer", rake) and not mentions("the drake is broken", rake)


def test_praise_gets_the_lowest_priority_so_its_card_can_show() -> None:
    praise = {k: v for k, v in COMPLETE.items() if k != "priority"} | {"category": "praise", "work_type": "no_change"}
    draft, errors = apply_update({}, praise, {"rake-agent"})
    assert not errors and draft["priority"] == "P4" and not missing(draft)
    bug, _ = apply_update({}, {k: v for k, v in COMPLETE.items() if k != "priority"}, {"rake-agent"})
    assert "priority" in missing(bug)  # anything else still waits for the impact question
    kept, _ = apply_update({}, {**praise, "priority": "P2"}, {"rake-agent"})
    assert kept["priority"] == "P2"  # a priority the model did set is left alone


def test_brief_lists_lose_the_models_own_numbers_and_bullets() -> None:
    """Review item 26: the pages number the steps, so "1. 1. Fill out a form" never shows again."""
    draft, errors = apply_update(
        {},
        {
            "steps": ["1. Fill out a form", "2) Press Submit", "Step 3: it resets", "- watch the spinner"],
            "open_gaps": ["• Which form", "a. Any error message"],
            "done_when": ["1.5 seconds is fine", "10:30 runs work"],
        },
        {"rake-agent"},
    )
    assert not errors
    assert draft["brief"]["steps"] == ["Fill out a form", "Press Submit", "it resets", "watch the spinner"]
    assert draft["brief"]["open_gaps"] == ["Which form", "Any error message"]
    assert draft["brief"]["done_when"] == ["1.5 seconds is fine", "10:30 runs work"]  # numbers in the text stay


def test_praise_card_shows_praise_alone_without_a_priority() -> None:
    praise, _ = apply_update({}, COMPLETE | {"category": "praise", "work_type": "no_change"}, {"rake-agent"})
    shown = card(praise, "RAKE Agent")
    assert (shown["work_type"], shown["category"], shown["priority"]) == ("", "Praise", "")
    bug, _ = apply_update({}, COMPLETE, {"rake-agent"})
    assert card(bug, "RAKE Agent")["priority"] == bug["priority"]


def test_draft_hash_changes_with_any_field() -> None:
    a, _ = apply_update({}, COMPLETE, {"rake-agent"})
    b, _ = apply_update(a, {"priority": "P1"}, {"rake-agent"})
    assert draft_hash(a) != draft_hash(b)


@dataclass
class FakeSession:
    id: str = "s1"
    events: list[Any] = field(default_factory=list)


@dataclass
class FakeContext:
    state: dict[str, Any] = field(default_factory=dict)
    session: FakeSession = field(default_factory=FakeSession)


@dataclass
class FakeTool:
    name: str


def guard(name: str, ctx: FakeContext, args: dict[str, Any] | None = None) -> dict[str, Any] | None:
    return callbacks.before_tool(FakeTool(name), args or {}, ctx)  # type: ignore[arg-type]


def test_submit_guards() -> None:
    ctx = FakeContext(state={"draft": dict(COMPLETE), "draft_confirmed": False})
    assert (guard("submit_feedback", ctx) or {}).get("error") == "not_confirmed"
    ctx.state["draft_confirmed"] = True
    assert (guard("submit_feedback", ctx, {"summary": "unseen"}) or {}).get("error") == "no_arguments"
    assert guard("submit_feedback", ctx) is None
    assert (guard("submit_feedback", ctx) or {}).get("error") == "limit_reached"


async def test_editing_the_draft_clears_the_confirmation() -> None:
    ctx = FakeContext(state={"draft": dict(COMPLETE), "draft_confirmed": True})
    await tools.update_draft(ctx, priority="P1")  # type: ignore[arg-type]
    assert ctx.state["draft_confirmed"] is False


async def test_submit_tool_refuses_without_confirm_even_if_the_guard_is_skipped() -> None:
    ctx = FakeContext(
        state={"draft": dict(COMPLETE), "draft_confirmed": False, "user": {"email": "a@redadair.example", "name": "A"}}
    )
    assert (await tools.submit_feedback(ctx))["error"] == "not_confirmed"  # type: ignore[arg-type]


# --- model credentials ---------------------------------------------------------------------


def test_choosing_gemini_exports_the_key_before_the_agent_is_built(monkeypatch: pytest.MonkeyPatch) -> None:
    # Locally .env reaches Settings only; embeddings make their own genai client from the environment.
    settings = get_settings().model_copy(update={"model_provider": "gemini", "google_api_key": "test-key"})
    monkeypatch.setattr(model, "get_settings", lambda: settings)
    monkeypatch.setenv("GOOGLE_API_KEY", "")
    monkeypatch.setenv("GOOGLE_GENAI_USE_ENTERPRISE", "true")
    agent_runtime.reset()
    try:
        assert agent_runtime.model_choice().provider == "gemini"
        assert os.environ["GOOGLE_API_KEY"] == "test-key"
        assert os.environ["GOOGLE_GENAI_USE_ENTERPRISE"] == "false"
    finally:
        agent_runtime.reset()


def test_a_summary_about_the_user_is_refused() -> None:
    draft, errors = apply_update({}, {"summary": "The user reported the system is slow.", "title": "Slow"}, set())
    assert "summary" not in draft and draft["title"] == "Slow"
    assert errors and "never 'The user" in errors[0]


def test_the_card_shows_what_happened_and_the_open_gaps_for_a_fix() -> None:
    draft = {
        "work_type": "fix",
        "category": "performance",
        "brief": {"actual": "Search takes a minute", "open_gaps": ["Since when"]},
    }
    shown = card(draft, "IT Asset")
    assert shown["what_happened"] == "Search takes a minute" and shown["gaps"] == ["Since when"]
    assert card({**draft, "work_type": "new_feature"}, "IT Asset")["what_happened"] == ""
    assert card({**draft, "category": "praise"}, "IT Asset")["gaps"] == []


def test_impact_chips_fit_a_fault_or_an_idea() -> None:
    fault = [c["label"] for c in impact_chips({"work_type": "fix", "category": "performance"})]
    idea = [c["label"] for c in impact_chips({"work_type": "new_feature", "category": "feature_request"})]
    assert "Can't work at all" in fault and not any("nice" in label.lower() for label in fault)
    assert "Nice to have" in idea and "Can't work at all" not in idea


def test_suggested_replies_are_short_and_promise_nothing_and_name_nobody() -> None:
    kept = quick_replies(
        [
            "Every time",
            "every time",
            "Since Monday",
            "We'll fix it today",
            "Ask Darlene",
            "Which screen?",
            "x" * 41,
            "Only on my laptop",
        ],
        ["Darlene Antonette Duran"],
    )
    assert kept == ["Every time", "Since Monday", "Only on my laptop"]
